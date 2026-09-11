import asyncio
from datetime import datetime, timezone
import importlib.metadata
import ipaddress
import json
from pathlib import Path
import shutil
import socket
import subprocess
import sys
import zipfile

import httpx
from bs4 import BeautifulSoup

from .store import confined_path


PACKAGES = {"pdf": "pypdf", "docx": "python-docx", "xlsx": "openpyxl",
            "html": "beautifulsoup4", "web_search": "ddgs"}
MAX_INPUT = 50 * 1024 * 1024
MAX_TEXT = 200_000


def environment_status():
    from .browser import installed_programs
    programs = installed_programs()
    packages = {}
    for feature, package in PACKAGES.items():
        try:
            packages[feature] = {"installed": True, "version": importlib.metadata.version(package)}
        except importlib.metadata.PackageNotFoundError:
            packages[feature] = {"installed": False}
    return {"checked_at": datetime.now(timezone.utc).isoformat(), "python": sys.version.split()[0],
            "packages": packages,
            "executables_on_path": {name: bool(shutil.which(name)) for name in ("uv", "ffmpeg", "tesseract", "soffice")},
            "document_python": sys.executable,
            "scope": "Service environment only; user project environments may differ. Installation is not a successful workload test.",
            "not_checked": ["network connectivity", "OCR or speech model weights", "GPU", "account authorization"],
            "office_license_required": False,
            "installed_programs": programs,
            "browser_integration": True, "browser_available": bool(programs["chrome"]),
            "browser_profile": "Separate per-user persistent profile; no everyday Chrome profile access",
            "calendar_integration": False}


def search_files(store, user_id, through_message, query=""):
    store.user_root(user_id)
    return [dict(row) for row in store.db.execute(
        "SELECT number,name,source FROM files WHERE user_id=? AND instr(lower(name),lower(?))>0 "
        "AND number NOT IN (SELECT number FROM deleted_files) AND ("
        "(source='original' AND EXISTS (SELECT 1 FROM attachments WHERE user_id=files.user_id "
        "AND path=files.path AND message_id<=? AND status='ready')) OR "
        "(source='outbound' AND EXISTS (SELECT 1 FROM outbox WHERE user_id=files.user_id "
        "AND content=files.path AND kind='file' AND message_id<=?))) ORDER BY number DESC LIMIT 30",
        (user_id, query[:200], through_message, through_message))]


def prepare_file(store, user_id, through_message, number):
    record = store.file_by_number(user_id, number)
    if record["source"] == "original":
        eligible = store.db.execute("SELECT 1 FROM attachments WHERE user_id=? AND path=? "
                                    "AND message_id<=? AND status='ready'",
                                    (user_id, record["path"], through_message)).fetchone()
        source = confined_path(store.user_root(user_id) / "originals", record["path"])
    else:
        eligible = store.db.execute("SELECT 1 FROM outbox WHERE user_id=? AND content=? "
                                    "AND kind='file' AND message_id<=?",
                                    (user_id, record["path"], through_message)).fetchone()
        source = confined_path(store.root / "outbound", str(confined_path(store.root, record["path"])))
    if not eligible or not source.is_file() or source.stat().st_size > MAX_INPUT:
        raise ValueError("File unavailable for this request")
    relative = f"retrieved/{number}/{record['name']}"
    target = confined_path(store.user_root(user_id) / "workspace", relative)
    target.parent.mkdir(parents=True, exist_ok=True)
    if not target.exists():
        shutil.copyfile(source, target)
    return {"number": number, "path": relative, "existing_working_copy_preserved": True}


def extract_document(path):
    path = Path(path)
    if not path.is_file() or path.stat().st_size > MAX_INPUT:
        raise ValueError("Document unavailable or larger than 50 MiB")
    suffix = path.suffix.lower()
    if suffix in (".docx", ".xlsx"):
        with zipfile.ZipFile(path) as archive:
            if len(archive.infolist()) > 10000 or sum(entry.file_size for entry in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("Expanded document too large")
    pieces = []
    size = 0
    limited = False

    def append(text):
        nonlocal size, limited
        remaining = MAX_TEXT - size
        pieces.append(text[:remaining])
        size += min(len(text), remaining)
        limited = limited or size >= MAX_TEXT
        return size >= MAX_TEXT

    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError("Encrypted PDF requires an unlocked working copy")
        for index, page in enumerate(reader.pages):
            if index >= 200:
                limited = True
                break
            if append(f"\n[Page {index + 1}]\n" + (page.extract_text() or "")):
                break
    elif suffix == ".docx":
        from docx import Document
        from docx.table import Table
        document = Document(path)
        for block in document.iter_inner_content():
            if isinstance(block, Table):
                text = "\n".join("\t".join(cell.text for cell in row.cells) for row in block.rows)
            else:
                text = block.text
            if append(text + "\n"):
                break
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            for sheet in workbook:
                if append(f"\n[Sheet {sheet.title}]\n"):
                    break
                for index, row in enumerate(sheet.iter_rows(max_col=min(sheet.max_column or 100, 100), values_only=True)):
                    if index >= 2000:
                        limited = True
                        break
                    if append("\t".join("" if value is None else str(value) for value in row) + "\n"):
                        break
                if (sheet.max_column or 0) > 100:
                    limited = True
                if size >= MAX_TEXT:
                    break
        finally:
            workbook.close()
    elif suffix in (".txt", ".md", ".csv", ".json", ".log", ".html", ".htm"):
        with path.open(encoding="utf-8-sig", errors="replace") as stream:
            text = stream.read(MAX_TEXT + 1)
        limited = len(text) > MAX_TEXT
        if suffix in (".html", ".htm"):
            text = html_text(text)
        append(text)
    else:
        return {"supported": False, "reason": "Use workspace scripts for other formats; OCR and speech models are not bundled."}
    return {"supported": True, "text": "".join(pieces), "extraction_limited": limited,
            "notes": "Text extraction only: no OCR, image interpretation, formula calculation or layout fidelity. PDF cap 200 pages; XLSX cap 2000 rows and 100 columns per sheet; text cap 200000 characters."}


async def read_document(workspace, relative, offset=0):
    if not 0 <= offset < MAX_TEXT:
        raise ValueError("Offset outside extraction range")
    path = confined_path(workspace, relative)
    process = await asyncio.create_subprocess_exec(
        sys.executable, "-m", "wechat_agent.foundation", str(path),
        stdout=subprocess.PIPE, stderr=subprocess.DEVNULL,
        **({"creationflags": subprocess.CREATE_NO_WINDOW} if sys.platform == "win32" else {}))
    try:
        output, _ = await asyncio.wait_for(process.communicate(), timeout=45)
    except BaseException:
        if process.returncode is None:
            process.kill()
        await process.wait()
        raise
    if process.returncode:
        return {"supported": False, "reason": "Parser failed; document may be encrypted, malformed or too large. Try a workspace script or an unlocked copy."}
    result = json.loads(output)
    text = result.pop("text", "")
    result.update(text=text[offset:offset + 12000], offset=offset,
                  next_offset=offset + 12000 if len(text) > offset + 12000 else None)
    return result


def html_text(content):
    soup = BeautifulSoup(content, "html.parser")
    for element in soup(["script", "style", "noscript", "nav", "footer"]):
        element.decompose()
    return soup.get_text("\n", strip=True)


async def public_address(url):
    parsed = httpx.URL(url)
    port = parsed.port or (443 if parsed.scheme == "https" else 80)
    if parsed.scheme not in ("http", "https") or not parsed.host or parsed.userinfo or port != (443 if parsed.scheme == "https" else 80):
        raise ValueError("Only public HTTP(S) URLs on standard ports are supported")
    addresses = await asyncio.get_running_loop().getaddrinfo(parsed.host, port, type=socket.SOCK_STREAM)
    if not addresses or any(not ipaddress.ip_address(item[4][0]).is_global for item in addresses):
        raise ValueError("Private or non-global address blocked")
    return parsed, addresses[0][4][0]


async def fetch_page(url):
    async with asyncio.timeout(30):
        async with httpx.AsyncClient(trust_env=False, follow_redirects=False, timeout=15) as client:
            for _ in range(5):
                parsed, address = await public_address(url)
                client.cookies.clear()
                async with client.stream("GET", parsed.copy_with(host=address),
                                         headers={"Host": parsed.netloc.decode("ascii"), "User-Agent": "WechatAgent/0.1"},
                                         extensions={"sni_hostname": parsed.host}) as response:
                    if response.is_redirect:
                        url = str(parsed.join(response.headers["location"]))
                        continue
                    response.raise_for_status()
                    kind = response.headers.get("content-type", "").split(";")[0].lower()
                    if kind not in ("text/html", "text/plain", "application/xhtml+xml"):
                        return {"url": str(parsed), "supported": False, "reason": "Only HTML and plain text; use workspace download tools for documents."}
                    content = bytearray()
                    async for chunk in response.aiter_bytes():
                        content.extend(chunk)
                        if len(content) > 2 * 1024 * 1024:
                            raise ValueError("Page larger than 2 MiB")
                    text = bytes(content).decode(response.encoding or "utf-8", errors="replace")
                    text = html_text(text) if kind != "text/plain" else text
                    return {"url": str(parsed), "text": text[:16000], "truncated": len(text) > 16000,
                            "source": "Untrusted public page; no JavaScript execution or login."}
    raise ValueError("Too many redirects")


async def web_search(query):
    query = query.strip()
    if not query or len(query) > 300:
        raise ValueError("Search query must be 1-300 characters")
    def search():
        from ddgs import DDGS
        results = DDGS(timeout=12).text(query, max_results=5)
        return [{"title": str(item.get("title", ""))[:300], "url": str(item.get("href", ""))[:2000],
                 "snippet": str(item.get("body", ""))[:1500]} for item in results[:5]]
    try:
        results = await asyncio.wait_for(asyncio.to_thread(search), timeout=20)
        return {"results": results, "source": "Public search snippets, not verified page contents"}
    except Exception:
        return {"results": [], "error": "Search provider unavailable or timed out. Do not invent results; try a supplied URL."}


if __name__ == "__main__":
    try:
        print(json.dumps(extract_document(sys.argv[1]), ensure_ascii=True))
    except Exception:
        raise SystemExit(1)