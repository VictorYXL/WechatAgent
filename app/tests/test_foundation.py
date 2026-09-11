import asyncio
import socket
from unittest.mock import AsyncMock, Mock

import httpx
import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from wechat_agent import foundation
from wechat_agent.store import Store


def test_file_search_and_preparation_are_owned_and_request_scoped(tmp_path):
    store = Store(tmp_path)
    try:
        user = store.user("bot", "user")
        other = store.user("bot", "other")
        first = store.ingest(user, "1", "", {})
        attachment = store.allocate_attachment(user, first, 0, "100% notes.txt", "file")
        store.save_attachment(attachment, b"owned content")
        current = store.ingest(user, "2", "read", {})
        future = store.ingest(user, "3", "", {})
        later = store.allocate_attachment(user, future, 0, "future.txt", "file")
        store.save_attachment(later, b"not yet")
        records = foundation.search_files(store, user, current, "%")
        assert len(records) == 1 and "path" not in records[0]
        number = records[0]["number"]
        assert len(foundation.search_files(store, user, current)) == 1
        assert not foundation.search_files(store, other, future)
        with pytest.raises(ValueError):
            foundation.prepare_file(store, other, current, number)
        with pytest.raises(ValueError):
            foundation.prepare_file(store, user, current, store.file_numbers(user)[-1])
        result = foundation.prepare_file(store, user, current, number)
        working = store.user_root(user) / "workspace" / result["path"]
        assert working.read_bytes() == b"owned content"
        working.write_bytes(b"edited")
        foundation.prepare_file(store, user, current, number)
        assert working.read_bytes() == b"edited"
        store.delete_files(user, [number])
        assert not foundation.search_files(store, user, current)
    finally:
        store.close()


def test_document_formats_and_pagination(tmp_path):
    document = Document()
    document.add_paragraph("Report heading")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Table value"
    document.save(tmp_path / "report.docx")
    workbook = Workbook()
    workbook.active.append(["Name", "Value"])
    workbook.active.append(["Total", "=1+2"])
    workbook.save(tmp_path / "report.xlsx")
    pdf = PdfWriter()
    page = pdf.add_blank_page(width=100, height=100)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 10 50 Td (PDF sample text) Tj ET")
    page[NameObject("/Contents")] = stream
    pdf.write(tmp_path / "empty.pdf")
    assert "Table value" in foundation.extract_document(tmp_path / "report.docx")["text"]
    assert "=1+2" in foundation.extract_document(tmp_path / "report.xlsx")["text"]
    assert "PDF sample text" in foundation.extract_document(tmp_path / "empty.pdf")["text"]
    (tmp_path / "long.txt").write_text("a" * 13000, encoding="utf-8")
    async def scenario():
        first = await foundation.read_document(tmp_path, "long.txt")
        assert len(first["text"]) == 12000 and first["next_offset"] == 12000
        last = await foundation.read_document(tmp_path, "long.txt", 12000)
        assert len(last["text"]) == 1000 and last["next_offset"] is None
        with pytest.raises(ValueError):
            await foundation.read_document(tmp_path, "../private.txt")
        (tmp_path / "broken.pdf").write_bytes(b"broken")
        assert not (await foundation.read_document(tmp_path, "broken.pdf"))["supported"]
    asyncio.run(scenario())


def test_environment_detection_does_not_claim_execution_or_accounts():
    result = foundation.environment_status()
    assert all(item["installed"] for item in result["packages"].values())
    assert "account authorization" in result["not_checked"]
    assert result["office_license_required"] is False
    assert not result["calendar_integration"] and not result["browser_integration"]


@pytest.mark.parametrize("failure", [TimeoutError, asyncio.CancelledError])
def test_document_timeout_or_cancel_kills_child(tmp_path, monkeypatch, failure):
    process = Mock(returncode=None)
    process.communicate = AsyncMock(side_effect=failure())
    process.wait = AsyncMock()
    monkeypatch.setattr(foundation.asyncio, "create_subprocess_exec", AsyncMock(return_value=process))
    with pytest.raises(failure):
        asyncio.run(foundation.read_document(tmp_path, "report.pdf"))
    process.kill.assert_called_once()
    process.wait.assert_awaited_once()


@pytest.mark.parametrize("url", ["file:///secret", "http://user:password@example.com", "http://example.com:8080"])
def test_fetch_rejects_non_public_url_forms(url):
    with pytest.raises(ValueError):
        asyncio.run(foundation.public_address(url))


@pytest.mark.parametrize("address", ["127.0.0.1", "10.0.0.1", "169.254.169.254", "::1", "192.168.1.1"])
def test_fetch_rejects_private_dns_answers(address, monkeypatch):
    async def scenario():
        async def resolve(*args, **kwargs):
            return [(socket.AF_INET, socket.SOCK_STREAM, 6, "", (address, 80))]
        monkeypatch.setattr(asyncio.get_running_loop(), "getaddrinfo", resolve)
        with pytest.raises(ValueError):
            await foundation.public_address("http://example.com")
    asyncio.run(scenario())


def test_fetch_pins_dns_and_rechecks_redirect(monkeypatch):
    client_class = httpx.AsyncClient
    seen = []
    async def resolve(url):
        if "internal" in url:
            raise ValueError("private")
        return httpx.URL(url), "93.184.216.34"
    def respond(request):
        seen.append(request)
        assert request.url.host == "93.184.216.34"
        assert request.headers["host"] == "example.com"
        assert request.extensions["sni_hostname"] == "example.com"
        if request.url.path == "/redirect":
            return httpx.Response(302, headers={"location": "http://internal/"})
        return httpx.Response(200, headers={"content-type": "text/html"},
                              text="<h1>Title</h1><script>secret()</script><p>Body</p>")
    monkeypatch.setattr(foundation, "public_address", resolve)
    monkeypatch.setattr(foundation.httpx, "AsyncClient", lambda **kwargs: client_class(transport=httpx.MockTransport(respond), **kwargs))
    result = asyncio.run(foundation.fetch_page("https://example.com/"))
    assert "Body" in result["text"] and "secret" not in result["text"]
    with pytest.raises(ValueError):
        asyncio.run(foundation.fetch_page("https://example.com/redirect"))
    assert len(seen) == 2


def test_public_dns_default_port_is_accepted(monkeypatch):
    async def scenario():
        async def resolve(host, port, **kwargs):
            assert host == "example.com" and port == 443
            return [(socket.AF_INET, socket.SOCK_STREAM, 6, "", ("93.184.216.34", port))]
        monkeypatch.setattr(asyncio.get_running_loop(), "getaddrinfo", resolve)
        parsed, address = await foundation.public_address("https://example.com")
        assert address == "93.184.216.34" and parsed.host == "example.com"
    asyncio.run(scenario())


def test_fetch_does_not_forward_response_cookies(monkeypatch):
    client_class = httpx.AsyncClient
    async def resolve(url):
        return httpx.URL(url), "93.184.216.34"
    def respond(request):
        assert "cookie" not in request.headers
        if request.url.path == "/":
            return httpx.Response(302, headers={"location": "/result", "set-cookie": "session=private; Path=/"})
        return httpx.Response(200, headers={"content-type": "text/plain"}, text="Public result")
    monkeypatch.setattr(foundation, "public_address", resolve)
    monkeypatch.setattr(foundation.httpx, "AsyncClient", lambda **kwargs: client_class(transport=httpx.MockTransport(respond), **kwargs))
    assert asyncio.run(foundation.fetch_page("https://example.com/"))["text"] == "Public result"


def test_search_returns_sources_and_explicit_failure(monkeypatch):
    import ddgs
    provider = Mock()
    provider.text.return_value = [{"title": "Title", "href": "https://example.com", "body": "Snippet"}]
    monkeypatch.setattr(ddgs, "DDGS", Mock(return_value=provider))
    result = asyncio.run(foundation.web_search("test query"))
    assert result["results"][0]["url"] == "https://example.com"
    provider.text.side_effect = RuntimeError("sensitive diagnostic")
    result = asyncio.run(foundation.web_search("test query"))
    assert result["results"] == [] and "sensitive" not in result["error"]